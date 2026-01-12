import os
from io import BytesIO
from typing import Dict, Any, List
from datetime import datetime

import fitz
from PIL import Image
from docx import Document
from docx2pdf import convert as docx2pdf_convert
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaFileUpload

from .config import (
    OUTPUT_LOCAL_DIR_DOC,
    OUTPUT_LOCAL_DIR_PDF,
    OUTPUT_LOCAL_DIR_JPEG,
    SHARED_DRIVE_ID,
    FILE_NAME_PATTERN,
    TEMPLATE_PATH,
    log,
)
from .formatters import fmt_number
from .template_engine import (
    build_mapping_for_owner,
    replace_placeholders_doc,
    add_assets_preserve_formatting,
)


def save_docx_locally(template_path: str, output_path: str, mapping: dict, items: list):
    """
    Args:
        template_path: path to template DOCX file
        output_path: path where the output DOCX should be saved
        mapping: dictionary of placeholder replacements
        items: list of asset items to add to the table
    """
    doc = Document(template_path)

    replace_placeholders_doc(doc, mapping)

    asset_table, header_idx = None, None
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            if any("Назва об’єкта" in (cell.text or "") for cell in row.cells):
                asset_table, header_idx = table, i
                break
        if asset_table:
            break

    if not asset_table:
        raise ValueError("Asset table not found")

    add_assets_preserve_formatting(asset_table, header_idx, items)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)


def upload_to_drive(drive_service, file_path: str, file_name: str) -> str:
    """Upload file to Google Drive shared drive.

    Args:
        drive_service: Google Drive API service instance
        file_path: Local path to file to upload
        file_name: Name for the file in Drive

    Returns:
        Google Drive file ID

    Raises:
        RuntimeError: If upload fails
    """
    try:
        media = MediaFileUpload(
            file_path,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        file_metadata = {"name": file_name, "parents": [SHARED_DRIVE_ID]}

        file = drive_service.files().create(
            body=file_metadata, media_body=media, supportsAllDrives=True
        ).execute()

        return file.get("id")

    except HttpError as e:
        raise RuntimeError(f"Drive upload failed: {e}")


def convert_to_pdf(docx_path: str) -> str:
    """Convert DOCX to PDF.

    Args:
        docx_path: Path to DOCX file

    Returns:
        Path to created PDF file

    Raises:
        RuntimeError: If conversion fails
    """
    pdf_path = docx_path.replace(".docx", ".pdf")
    pdf_dir = os.path.dirname(pdf_path).replace(OUTPUT_LOCAL_DIR_DOC, OUTPUT_LOCAL_DIR_PDF)
    pdf_path = os.path.join(pdf_dir, os.path.basename(pdf_path))

    os.makedirs(pdf_dir, exist_ok=True)

    try:
        docx2pdf_convert(docx_path, pdf_path)
        return pdf_path
    except Exception as e:
        raise RuntimeError(f"PDF conversion failed: {e}")


def convert_to_jpeg(pdf_path: str) -> str:
    """Convert first page of PDF to JPEG.

    Args:
        pdf_path: Path to PDF file

    Returns:
        Path to created JPEG file

    Raises:
        RuntimeError: If conversion fails
    """
    jpeg_path = pdf_path.replace(".pdf", ".jpg")
    jpeg_dir = os.path.dirname(jpeg_path).replace(OUTPUT_LOCAL_DIR_PDF, OUTPUT_LOCAL_DIR_JPEG)
    jpeg_path = os.path.join(jpeg_dir, os.path.basename(jpeg_path))

    os.makedirs(jpeg_dir, exist_ok=True)

    try:
        doc = fitz.open(pdf_path)
        page = doc[0]

        # Calculate zoom to fit longest side to 1280px
        rect = page.rect
        zoom_x = 1280 / max(rect.width, rect.height)
        mat = fitz.Matrix(zoom_x, zoom_x)

        pix = page.get_pixmap(matrix=mat, alpha=False)
        img_data = pix.tobytes("ppm")

        img = Image.open(BytesIO(img_data))

        img.save(jpeg_path, "JPEG", quality=100, optimize=False)
        doc.close()

        return jpeg_path

    except Exception as e:
        raise RuntimeError(f"JPEG conversion failed: {e}")


def create_act_docs_local(per_owner: Dict[str, Any], drive_service) -> List[Dict[str, Any]]:
    """Save all documents locally and upload to Google Drive.

    Args:
        per_owner: Dictionary mapping owner codes to their data
        drive_service: Google Drive API service instance

    Returns:
        List of created document info dictionaries
    """
    created = []
    upload_failed = []

    for code, data in per_owner.items():
        if not data["items"]:
            log.info(f"Owner {code} has no items; skipping.")
            continue

        dept = data["dept"]
        date_str = datetime.now().strftime("%Y %m %d")
        file_name = FILE_NAME_PATTERN.format(date=date_str, deptname=dept.get("code"))
        mapping = build_mapping_for_owner(data, dept)

        try:
            # Create DOCX
            docx_path = os.path.join(OUTPUT_LOCAL_DIR_DOC, f"{file_name}.docx")
            save_docx_locally(
                template_path=TEMPLATE_PATH,
                output_path=docx_path,
                mapping=mapping,
                items=data["items"],
            )

            doc_info = {
                "docx_path": docx_path,
                "name": file_name,
                "items": len(data["items"]),
                "sum": data["tot_sum"],
            }

            # Upload to Google Drive
            try:
                drive_file_id = upload_to_drive(drive_service, docx_path, f"{file_name}.docx")
                doc_info["drive_file_id"] = drive_file_id
                log.info(
                    f'Created and uploaded "{file_name}.docx" (ID: {drive_file_id}) - '
                    f'items={len(data["items"])} - sum={fmt_number(data["tot_sum"])}'
                )
            except Exception as e:
                log.warning(f"Drive upload failed for {code}: {e}")
                upload_failed.append(code)
                log.info(
                    f'Created local "{docx_path}" (upload failed) - '
                    f'items={len(data["items"])} - sum={fmt_number(data["tot_sum"])}'
                )

            created.append(doc_info)

            # Create PDF and JPEG
            try:
                pdf_path = convert_to_pdf(docx_path)
                doc_info["pdf_path"] = pdf_path

                try:
                    jpeg_path = convert_to_jpeg(pdf_path)
                    doc_info["jpeg_path"] = jpeg_path
                    log.info(
                        f'Created docs "{docx_path}" + PDF + JPEG - '
                        f'items={len(data["items"])} - sum={fmt_number(data["tot_sum"])}'
                    )
                except Exception as jpeg_err:
                    log.warning(f"JPEG conversion failed for {code}: {jpeg_err}")
                    log.info(
                        f'Created docs "{docx_path}" + PDF - '
                        f'items={len(data["items"])} - sum={fmt_number(data["tot_sum"])}'
                    )

            except Exception as e:
                log.warning(f"PDF conversion failed for {code}: {e}")
                log.info(
                    f'Created doc "{docx_path}" (PDF/JPEG skipped) - '
                    f'items={len(data["items"])} - sum={fmt_number(data["tot_sum"])}'
                )

        except Exception as e:
            log.error(f"Document creation failed for {code}: {e}")
            continue

    if upload_failed:
        log.info(f"Drive upload failed for: {', '.join(upload_failed)}")

    return created
